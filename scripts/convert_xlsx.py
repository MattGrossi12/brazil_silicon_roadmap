<<<<<<< HEAD
﻿#!/usr/bin/env python3
"""Convert the project spreadsheets into website JSON data."""
=======
#!/usr/bin/env python3
"""Convert the bilingual Brazil Silicon Roadmap workbooks into website JSON."""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


CLUSTERS = [
    ("manaus", "Manaus", "Manaus", -3.1190, -60.0217, "AM"),
    ("brasilia", "Brasília", "Brasília", -15.7939, -47.8828, "DF"),
    ("campina-grande", "Campina Grande", "Campina Grande", -7.2307, -35.8817, "PB"),
    ("salvador", "Salvador", "Salvador", -12.9777, -38.5016, "BA"),
    ("minas-interior", "Minas Gerais State Interior", "Interior de Minas Gerais", -22.2537, -45.7030, "MG"),
    ("belo-horizonte", "Belo Horizonte and Surrounding Area", "Belo Horizonte e proximidades", -19.9167, -43.9345, "MG"),
    ("campinas", "Campinas and Surrounding Area", "Campinas e proximidades", -22.9056, -47.0608, "SP"),
    ("florianopolis", "Florianópolis", "Florianópolis", -27.5949, -48.5482, "SC"),
    ("porto-alegre", "Porto Alegre and Surrounding Area", "Porto Alegre e proximidades", -30.0346, -51.2177, "RS"),
    ("sao-paulo-interior", "São Paulo State Interior", "Interior de São Paulo", -22.0174, -47.8860, "SP"),
    ("curitiba", "Curitiba and Surrounding Area", "Curitiba e proximidades", -25.4284, -49.2733, "PR"),
    ("rio-grande-do-sul-interior", "Rio Grande do Sul State Interior", "Interior do Rio Grande do Sul", -29.6990, -54.7730, "RS"),
    ("sao-jose-dos-campos", "São José dos Campos and Surrounding Area", "São José dos Campos e proximidades", -23.2237, -45.9009, "SP"),
]

CATEGORY_LABELS = {
    "brazilian-private": {"en": "Brazilian private company", "pt": "Empresa privada nacional"},
    "international-private": {"en": "International private company", "pt": "Empresa privada internacional"},
    "private-research": {"en": "Private research center", "pt": "Centro de pesquisa privado"},
    "public-research": {"en": "Public research center", "pt": "Centro de pesquisa público"},
    "state-owned": {"en": "State-owned company", "pt": "Empresa estatal"},
}


def normalized(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return re.sub(
        r"[^a-z0-9]+",
        " ",
        text.encode("ascii", "ignore").decode().lower(),
    ).strip()


CLUSTER_LOOKUP = {
    normalized(label): cluster_id
    for cluster_id, english, portuguese, *_ in CLUSTERS
    for label in (english, portuguese)
}


def read_info(path: Path) -> dict[str, dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = {}
    for number, company, activity, organization_type, source in workbook[
        "info"
    ].iter_rows(min_row=2, max_col=5, values_only=True):
        if company:
            rows[str(company).strip()] = {
                "number": int(number),
                "activity": str(activity or "").strip(),
                "organizationType": str(organization_type or "").strip(),
                "source": str(source or "").strip(),
            }
    workbook.close()
    return rows


def read_locations(path: Path) -> list[tuple[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_name = "grouped" if "grouped" in workbook.sheetnames else "agrouped"
    current_cluster = None
    entries = []
    for first, second in workbook[sheet_name].iter_rows(
        min_row=2,
        max_col=2,
        values_only=True,
    ):
        if second is None and first is not None and not isinstance(first, (int, float)):
            current_cluster = CLUSTER_LOOKUP.get(normalized(first))
            if current_cluster is None:
                raise ValueError(f"Unknown location heading: {first}")
        elif second and current_cluster:
            entries.append((str(second).strip(), current_cluster))
    workbook.close()
    return entries


def category_for(organization_type: str) -> str:
    value = normalized(organization_type)
    if "state owned" in value or "estatal" in value:
        return "state-owned"
    if (
        "public research" in value
        or "pesquisas publica" in value
        or "pesquisa publico" in value
    ):
        return "public-research"
    if "research" in value or "pesquisa" in value:
        return "private-research"
    if "international" in value or "internacional" in value:
        return "international-private"
    return "brazilian-private"


def build_payload(english_path: Path, portuguese_path: Path) -> dict[str, object]:
    english = read_info(english_path)
    portuguese = read_info(portuguese_path)
    entries = read_locations(english_path)
    locations_by_company = defaultdict(list)
    companies_by_cluster = defaultdict(list)

    for company, cluster_id in entries:
        companies_by_cluster[cluster_id].append(company)
        if cluster_id not in locations_by_company[company]:
            locations_by_company[company].append(cluster_id)

    companies = []
    for company, row in sorted(english.items(), key=lambda item: item[1]["number"]):
        translated = portuguese.get(company, row)
        companies.append(
            {
                "id": normalized(company).replace(" ", "-"),
                "number": row["number"],
                "name": company,
                "activity": {
                    "en": row["activity"],
                    "pt": translated["activity"],
                },
                "organizationType": {
                    "en": row["organizationType"],
                    "pt": translated["organizationType"],
                },
                "category": category_for(str(row["organizationType"])),
                "sources": [
                    item.strip()
                    for item in str(row["source"]).split("|")
                    if item.strip()
                ],
                "locations": locations_by_company[company],
            }
        )

    clusters = []
    for cluster_id, english_label, portuguese_label, latitude, longitude, state in CLUSTERS:
        names = companies_by_cluster[cluster_id]
        clusters.append(
            {
                "id": cluster_id,
                "label": {"en": english_label, "pt": portuguese_label},
                "latitude": latitude,
                "longitude": longitude,
                "state": state,
                "companies": names,
                "count": len(names),
            }
        )

    return {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "source": "https://github.com/MattGrossi12/brazil_silicon_roadmap",
        "metadata": {
            "organizations": len(companies),
            "locationEntries": len(entries),
            "clusters": len(clusters),
        },
        "categoryLabels": CATEGORY_LABELS,
        "companies": companies,
        "clusters": clusters,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--english", type=Path, required=True)
    parser.add_argument("--portuguese", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    payload = build_payload(args.english, args.portuguese)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    metadata = payload["metadata"]
    print(
        f"Wrote {metadata['organizations']} organizations, "
        f"{metadata['locationEntries']} location entries, and "
        f"{metadata['clusters']} clusters to {args.output}"
    )


if __name__ == "__main__":
    main()
>>>>>>> feature/interactive-map
